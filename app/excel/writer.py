from __future__ import annotations

from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.logger import get_logger

logger = get_logger("excel_writer")

# ── Colours ───────────────────────────────────────────────────────────────────
_HDR_BG  = "1F3864"   # dark navy  — title row
_HDR_FG  = "FFFFFF"
_EMP_BG  = "2E75B6"   # blue       — employee name header (spans 2 cols)
_EMP_FG  = "FFFFFF"
_SUB_BG  = "4472C4"   # mid-blue   — Tasks / Status sub-headers
_SUB_FG  = "FFFFFF"
_DATE_BG = "1F3864"   # dark navy  — date cells
_DATE_FG = "FFFFFF"
_ROW_A   = "EBF3FB"   # light blue — alternating
_ROW_B   = "FFFFFF"

_STATUS_CLR = {
    "completed":   ("C6EFCE", "375623"),   # green
    "in_progress": ("FFEB9C", "9C5700"),   # amber
    "open":        ("DDEBF7", "1F3864"),   # light blue
    "blocked":     ("FFC7CE", "9C0006"),   # red
    "review":      ("E2EFDA", "375623"),   # pale green
}

_STATUS_LABEL = {
    "completed":   "✅ Completed",
    "in_progress": "🔄 In Progress",
    "open":        "📋 Open",
    "blocked":     "🚫 Blocked",
    "review":      "🔍 Review",
}

# Priority order for aggregating multiple tasks → one status per cell
_STATUS_ORDER = ["in_progress", "open", "review", "blocked", "completed"]

_THIN   = Side(style="thin",   color="B0BEC5")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color="000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Calibri")


def _normalize_day(raw_date: str) -> str:
    """Return YYYY-MM-DD from any reasonable date string."""
    raw = str(raw_date).strip()
    if " " in raw:
        raw = raw.split(" ")[0]
    if "T" in raw:
        raw = raw.split("T")[0]
    day = raw[:10]
    parts = day.split("-")
    if len(parts) == 3 and len(parts[0]) == 2 and len(parts[2]) == 4:
        return f"{parts[2]}-{parts[1]}-{parts[0]}"
    return day


def _day_label(iso: str) -> str:
    """YYYY-MM-DD → '08 Sep 2026'"""
    try:
        return date.fromisoformat(iso).strftime("%d %b %Y")
    except Exception:
        return iso


def _dominant_status(statuses: list[str]) -> str:
    """Pick the 'most active' status from a list (in_progress > open > review > blocked > completed)."""
    for s in _STATUS_ORDER:
        if s in statuses:
            return s
    return statuses[0] if statuses else "open"


class ExcelWriter:
    def __init__(self, file_path: str) -> None:
        self.file_path = Path(file_path).expanduser()
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

    # ── Public entry point ────────────────────────────────────────────────────
    def write_tasks(self, rows: list[dict]) -> None:
        """
        Write tasks to Excel.
        Strategy:
        - Load existing workbook (if any)
        - Read ALL existing rows from the sheet already there
        - Merge with incoming rows: new dates OVERWRITE old same-date data
        - Old dates NOT in the incoming rows are PRESERVED as-is
        - Rebuild the full sheet with merged data
        """
        if not rows:
            logger.info("No task rows to write to Excel.")
            return

        # ── 1. Parse incoming rows into data dict ─────────────────────────────
        incoming: dict[str, dict[str, dict]] = defaultdict(
            lambda: defaultdict(lambda: {"tasks": [], "statuses": []})
        )
        for r in rows:
            day      = _normalize_day(r.get("date") or r.get("created_at") or date.today().isoformat())
            assignee = str(r.get("assignee") or "Unknown").strip()
            task     = str(r.get("task") or "").strip()
            status   = str(r.get("status") or "open").lower().strip()
            for line in task.splitlines():
                line = line.strip().lstrip("•").strip()
                if line:
                    incoming[day][assignee]["tasks"].append(line)
            if status:
                incoming[day][assignee]["statuses"].append(status)

        incoming_dates = set(incoming.keys())

        # ── 2. Load existing data from workbook (preserve old dates) ──────────
        existing: dict[str, dict[str, dict]] = defaultdict(
            lambda: defaultdict(lambda: {"tasks": [], "statuses": []})
        )
        target = self.file_path
        if target.exists():
            try:
                wb_old = load_workbook(target)
                if "Task Summary" in wb_old.sheetnames:
                    ws_old = wb_old["Task Summary"]
                    existing = self._read_existing_data(ws_old, incoming_dates)
                    logger.info("Preserved %d old date rows from existing Excel.", len(existing))
            except Exception as e:
                logger.warning("Could not read existing workbook: %s — starting fresh.", e)

        # ── 3. Merge: existing (old dates) + incoming (new dates) ─────────────
        # incoming dates always win (overwrite); old dates are kept
        data: dict[str, dict[str, dict]] = defaultdict(
            lambda: defaultdict(lambda: {"tasks": [], "statuses": []})
        )
        for day, employees in existing.items():
            for emp, d in employees.items():
                data[day][emp]["tasks"]    = list(d["tasks"])
                data[day][emp]["statuses"] = list(d["statuses"])
        for day, employees in incoming.items():
            for emp, d in employees.items():
                data[day][emp]["tasks"]    = list(d["tasks"])
                data[day][emp]["statuses"] = list(d["statuses"])

        # ── 4. Write merged data ──────────────────────────────────────────────
        temp = target.with_suffix(".tmp.xlsx")
        wb   = Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        ws = wb.create_sheet(title="Task Summary", index=0)
        self._write_summary_sheet(ws, data)
        wb.active = wb["Task Summary"]

        wb.save(temp)
        temp.replace(target)
        logger.info("Excel workbook updated: %s (%d dates total)", target, len(data))

    # ── Read existing sheet back into data dict ───────────────────────────────
    def _read_existing_data(
        self,
        ws,
        skip_dates: set[str],
    ) -> dict[str, dict[str, dict]]:
        """
        Parse an existing Task Summary sheet back into data dict.
        Skips any dates in skip_dates (those will be replaced by incoming data).
        Reads row structure: col A = date label, then pairs (tasks, status) per employee.
        Row 1 = title, Row 2 = employee headers, Row 3 = sub-headers, Row 4+ = data.
        """
        data: dict[str, dict[str, dict]] = defaultdict(
            lambda: defaultdict(lambda: {"tasks": [], "statuses": []})
        )

        # Read employee names from row 2 (columns B, D, F, … — merged, so first col of each pair)
        employees: list[str] = []
        max_col = ws.max_column
        col = 2
        while col <= max_col:
            cell_val = ws.cell(row=2, column=col).value
            if cell_val and str(cell_val).strip():
                employees.append(str(cell_val).strip())
            col += 2   # each employee takes 2 cols

        if not employees:
            return data

        # Read data rows (row 4 onward)
        for row_idx in range(4, ws.max_row + 1):
            date_cell = ws.cell(row=row_idx, column=1).value
            if not date_cell:
                continue
            # Convert "08 Sep 2026" back to "2026-09-08"
            try:
                from datetime import datetime as _dt
                iso = _dt.strptime(str(date_cell).strip(), "%d %b %Y").date().isoformat()
            except Exception:
                iso = str(date_cell).strip()

            # Skip dates that will be overwritten by incoming data
            if iso in skip_dates:
                continue

            for i, emp in enumerate(employees):
                tasks_col  = 2 + i * 2
                status_col = tasks_col + 1
                tasks_cell  = ws.cell(row=row_idx, column=tasks_col).value or ""
                status_cell = ws.cell(row=row_idx, column=status_col).value or ""

                # Parse bullet points back into individual tasks
                for line in str(tasks_cell).splitlines():
                    line = line.strip().lstrip("•").strip()
                    if line:
                        data[iso][emp]["tasks"].append(line)

                # Reverse-map status label back to key
                status_reverse = {v: k for k, v in _STATUS_LABEL.items()}
                raw_status = str(status_cell).strip()
                status_key = status_reverse.get(raw_status, "open")
                if data[iso][emp]["tasks"]:
                    data[iso][emp]["statuses"].append(status_key)

        return data

    # ── Sheet builder ─────────────────────────────────────────────────────────
    def _write_summary_sheet(self, ws, data: dict) -> None:
        """
        Layout
        ──────
        Row 1 : Title  (merged across all columns)
        Row 2 : Date | ←──── Ayan ────→ | ←── Farah ──→ | ...
                       Tasks   Status      Tasks   Status
        Row 3 : 2 sub-header cols per employee
        Row 4+: data rows
        """
        sorted_dates  = sorted(data.keys())
        all_employees = sorted({emp for day in data.values() for emp in day})
        n_emp         = len(all_employees)
        total_cols    = 1 + n_emp * 2        # date col + (tasks + status) × employees

        # ── Row 1: Title ──────────────────────────────────────────────────────
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
        c = ws.cell(row=1, column=1, value="TenBit Daily Task Report")
        c.fill      = _fill(_HDR_BG)
        c.font      = _font(bold=True, color=_HDR_FG, size=14)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # ── Row 2: Date header + merged employee name headers ─────────────────
        # Date header (spans rows 2-3)
        ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)
        c = ws.cell(row=2, column=1, value="Date")
        c.fill      = _fill(_DATE_BG)
        c.font      = _font(bold=True, color=_DATE_FG, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _BORDER

        for i, emp in enumerate(all_employees):
            tasks_col  = 2 + i * 2       # e.g. col 2, 4, 6 …
            status_col = tasks_col + 1   # e.g. col 3, 5, 7 …

            # Employee name spans both sub-columns
            ws.merge_cells(start_row=2, start_column=tasks_col,
                           end_row=2,   end_column=status_col)
            c = ws.cell(row=2, column=tasks_col, value=emp)
            c.fill      = _fill(_EMP_BG)
            c.font      = _font(bold=True, color=_EMP_FG, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = _BORDER

            # Sub-headers row 3
            for col, label in [(tasks_col, "Tasks"), (status_col, "Status")]:
                c = ws.cell(row=3, column=col, value=label)
                c.fill      = _fill(_SUB_BG)
                c.font      = _font(bold=True, color=_SUB_FG, size=9)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border    = _BORDER

        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 16

        # ── Column widths ─────────────────────────────────────────────────────
        ws.column_dimensions["A"].width = 13   # Date
        for i in range(n_emp):
            tasks_col  = 2 + i * 2
            status_col = tasks_col + 1
            ws.column_dimensions[get_column_letter(tasks_col)].width  = 32
            ws.column_dimensions[get_column_letter(status_col)].width = 14

        # ── Data rows (start row 4) ───────────────────────────────────────────
        for row_idx, day in enumerate(sorted_dates, start=4):
            bg = _ROW_A if row_idx % 2 == 0 else _ROW_B

            # Date cell
            c = ws.cell(row=row_idx, column=1, value=_day_label(day))
            c.fill      = _fill(_DATE_BG)
            c.font      = _font(bold=True, color=_DATE_FG, size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = _BORDER

            max_lines = 1
            for i, emp in enumerate(all_employees):
                tasks_col  = 2 + i * 2
                status_col = tasks_col + 1

                emp_data = data[day].get(emp, {"tasks": [], "statuses": []})
                tasks    = emp_data["tasks"]
                statuses = emp_data["statuses"]

                # ── Tasks cell ────────────────────────────────────────────────
                cell_text = "\n".join(f"• {t}" for t in tasks) if tasks else ""
                max_lines = max(max_lines, len(tasks))

                c = ws.cell(row=row_idx, column=tasks_col, value=cell_text)
                c.fill      = _fill(bg)
                c.font      = _font(size=10)
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                c.border    = _BORDER

                # ── Status cell ───────────────────────────────────────────────
                if tasks:
                    status_key = _dominant_status(statuses) if statuses else "open"
                    s_bg, s_fg = _STATUS_CLR.get(status_key, ("E0E0E0", "333333"))
                    s_label    = _STATUS_LABEL.get(status_key, status_key.title())
                else:
                    s_bg, s_fg = ("F5F5F5", "AAAAAA")
                    s_label    = "—"

                c = ws.cell(row=row_idx, column=status_col, value=s_label)
                c.fill      = _fill(s_bg)
                c.font      = _font(bold=True, color=s_fg, size=9)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                c.border    = _BORDER

            ws.row_dimensions[row_idx].height = max(18, min(max_lines * 16, 250))

        # Freeze: lock Date column + both header rows while scrolling
        ws.freeze_panes = "B4"
